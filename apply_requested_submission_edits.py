from __future__ import annotations

from copy import deepcopy
import os
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
MAIN_DOCX = ROOT / "JPS_manuscript_draft_v2.docx"
MAIN_MD = ROOT / "JPS_manuscript_draft_v2.md"
SUPP_DOCX = ROOT / "Supplementary_Material.docx"
SUPP_MD = ROOT / "Supplementary_Material.md"
CAPTIONS_DOCX = ROOT / "Figure_captions.docx"

PHONE = "Tel: +86 18995563848."
S3_LABEL = "Supplementary Table S3."
S3_BODY = "Timing of unplanned reoperation (n=36)."
S2_LABEL = "Supplementary Figure S2."
S2_SHORT_BODY = (
    "Cause-specific rate of unplanned early reoperation by index surgical approach."
)
S2_EDITORIAL_NOTE = (
    " (graphical form of Table 2; moved from the main text to keep the combined "
    "table-and-figure count within the journal limit)"
)
S2_FILE_NOTE = " (File: FigureS2_cause_by_approach.png / .pdf)"
S2_FILE_NOTE_MARKDOWN = " *(File: FigureS2_cause_by_approach.png / .pdf)*"


def atomic_save_docx(doc: Document, path: Path) -> None:
    temp_path = path.with_name(f"{path.stem}.__editing__.docx")
    doc.save(temp_path)
    Document(temp_path)
    os.replace(temp_path, path)


def atomic_save_text(path: Path, text: str) -> None:
    temp_path = path.with_name(f"{path.name}.__editing__")
    temp_path.write_text(text, encoding="utf-8", newline="\n")
    os.replace(temp_path, path)


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one paragraph starting with {prefix!r}; found {len(matches)}"
        )
    return matches[0]


def set_labeled_paragraph(paragraph: Paragraph, label: str, body: str) -> None:
    paragraph.clear()
    label_run = paragraph.add_run(label)
    label_run.bold = True
    paragraph.add_run(f" {body}")


def insert_labeled_paragraph_after(
    template: Paragraph, label: str, body: str
) -> Paragraph:
    new_element = deepcopy(template._p)
    for child in list(new_element):
        if child.tag != qn("w:pPr"):
            new_element.remove(child)
    template._p.addnext(new_element)
    paragraph = Paragraph(new_element, template._parent)
    set_labeled_paragraph(paragraph, label, body)
    return paragraph


def update_main_markdown() -> None:
    lines = MAIN_MD.read_text(encoding="utf-8").splitlines()

    corresponding = [
        i for i, line in enumerate(lines) if line.startswith("**\\*Corresponding author:**")
    ]
    if len(corresponding) != 1:
        raise RuntimeError("Could not uniquely identify the corresponding-author line.")
    idx = corresponding[0]
    if "Tel:" not in lines[idx]:
        lines[idx] = f"{lines[idx]} {PHONE}"

    s2_indices = [
        i
        for i, line in enumerate(lines)
        if line.startswith("- **Supplementary Table S2.**")
    ]
    if len(s2_indices) != 1:
        raise RuntimeError("Could not uniquely identify Supplementary Table S2.")
    s2_idx = s2_indices[0]
    if not any(line.startswith("- **Supplementary Table S3.**") for line in lines):
        lines.insert(s2_idx + 1, f"- **{S3_LABEL}** {S3_BODY}")

    fig_s1_indices = [
        i
        for i, line in enumerate(lines)
        if line.startswith("- **Supplementary Figure S1.**")
    ]
    if len(fig_s1_indices) != 1:
        raise RuntimeError("Could not uniquely identify Supplementary Figure S1.")
    fig_s1_idx = fig_s1_indices[0]
    if not any(line.startswith("- **Supplementary Figure S2.**") for line in lines):
        lines.insert(fig_s1_idx + 1, f"- **{S2_LABEL}** {S2_SHORT_BODY}")

    atomic_save_text(MAIN_MD, "\n".join(lines) + "\n")


def update_main_docx() -> None:
    doc = Document(MAIN_DOCX)

    corresponding = find_paragraph(doc, "*Corresponding author:")
    if "Tel:" not in corresponding.text:
        if not corresponding.runs:
            raise RuntimeError("Corresponding-author paragraph has no runs.")
        corresponding.runs[-1].text = f"{corresponding.runs[-1].text} {PHONE}"

    s2 = find_paragraph(doc, "Supplementary Table S2.")
    if not any(p.text.startswith(S3_LABEL) for p in doc.paragraphs):
        insert_labeled_paragraph_after(s2, S3_LABEL, S3_BODY)

    figure_s1 = find_paragraph(doc, "Supplementary Figure S1.")
    if not any(p.text.startswith(S2_LABEL) for p in doc.paragraphs):
        insert_labeled_paragraph_after(figure_s1, S2_LABEL, S2_SHORT_BODY)

    atomic_save_docx(doc, MAIN_DOCX)


def clean_s2_caption(text: str) -> str:
    updated = text.replace(S2_EDITORIAL_NOTE, "")
    updated = updated.replace(S2_FILE_NOTE, "")
    updated = updated.replace(S2_FILE_NOTE_MARKDOWN, "")
    return updated


def update_supplement_markdown() -> None:
    lines = SUPP_MD.read_text(encoding="utf-8").splitlines()
    matches = [
        i
        for i, line in enumerate(lines)
        if line.startswith("**Supplementary Figure S2.** Cause-specific rate")
    ]
    if len(matches) != 1:
        raise RuntimeError("Could not uniquely identify the Supplementary Figure S2 caption.")
    idx = matches[0]
    lines[idx] = clean_s2_caption(lines[idx])
    atomic_save_text(SUPP_MD, "\n".join(lines) + "\n")


def update_supplement_docx() -> None:
    doc = Document(SUPP_DOCX)
    caption = [
        p
        for p in doc.paragraphs
        if p.text.startswith("Supplementary Figure S2. Cause-specific rate")
    ]
    if len(caption) != 1:
        raise RuntimeError(
            "Could not uniquely identify the Supplementary Figure S2 DOCX caption."
        )
    cleaned = clean_s2_caption(caption[0].text)
    body = cleaned[len(S2_LABEL) :].lstrip()
    set_labeled_paragraph(caption[0], S2_LABEL, body)
    atomic_save_docx(doc, SUPP_DOCX)


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_caption(doc: Document, label: str, body: str) -> None:
    paragraph = doc.add_paragraph(style="Caption Text")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, 12, bold=True)
    body_run = paragraph.add_run(f" {body}")
    set_run_font(body_run, 12)


def split_caption(text: str, label: str) -> str:
    if not text.startswith(label):
        raise RuntimeError(f"Caption does not start with {label!r}.")
    return text[len(label) :].lstrip()


def create_figure_captions_docx() -> None:
    main = Document(MAIN_DOCX)
    supplement = Document(SUPP_DOCX)

    main_figure_1 = find_paragraph(main, "Figure 1.")
    main_figure_2 = find_paragraph(main, "Figure 2.")
    supp_figure_s1 = find_paragraph(supplement, "Supplementary Figure S1.")
    supp_figure_s2 = [
        p
        for p in supplement.paragraphs
        if p.text.startswith("Supplementary Figure S2. Cause-specific rate")
    ]
    if len(supp_figure_s2) != 1:
        raise RuntimeError("Could not uniquely retrieve Supplementary Figure S2.")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    if "Caption Text" not in [style.name for style in doc.styles]:
        caption_style = doc.styles.add_style("Caption Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption_style = doc.styles["Caption Text"]
    caption_style.base_style = normal
    caption_style.font.name = "Times New Roman"
    caption_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption_style.font.size = Pt(12)
    caption_style.paragraph_format.space_before = Pt(0)
    caption_style.paragraph_format.space_after = Pt(10)
    caption_style.paragraph_format.line_spacing = 1.10
    caption_style.paragraph_format.keep_together = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("Figure captions")
    set_run_font(title_run, 16, bold=True)

    manuscript = doc.add_paragraph()
    manuscript.paragraph_format.space_after = Pt(12)
    manuscript_label = manuscript.add_run("Manuscript title: ")
    set_run_font(manuscript_label, 11, bold=True)
    manuscript_title = manuscript.add_run(
        "Why Children Return to Theater After a Ladd Procedure: The Cause of Early "
        "Unplanned Reoperation Differs by Age, and by Surgical Approach in Neonates"
    )
    set_run_font(manuscript_title, 11)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run("Main figures")
    set_run_font(heading_run, 12, bold=True)

    add_caption(doc, "Figure 1.", split_caption(main_figure_1.text, "Figure 1."))
    add_caption(doc, "Figure 2.", split_caption(main_figure_2.text, "Figure 2."))

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run("Supplementary figures")
    set_run_font(heading_run, 12, bold=True)

    add_caption(
        doc,
        "Supplementary Figure S1.",
        split_caption(supp_figure_s1.text, "Supplementary Figure S1."),
    )
    add_caption(
        doc,
        "Supplementary Figure S2.",
        split_caption(supp_figure_s2[0].text, "Supplementary Figure S2."),
    )

    doc.core_properties.title = "Figure captions"
    doc.core_properties.subject = "Journal of Pediatric Surgery submission"
    doc.core_properties.author = "Xin Wang"
    atomic_save_docx(doc, CAPTIONS_DOCX)


def validate_results() -> None:
    main = Document(MAIN_DOCX)
    supplement = Document(SUPP_DOCX)
    captions = Document(CAPTIONS_DOCX)

    assert PHONE in find_paragraph(main, "*Corresponding author:").text
    main_text = "\n".join(p.text for p in main.paragraphs)
    assert f"{S3_LABEL} {S3_BODY}" in main_text
    assert f"{S2_LABEL} {S2_SHORT_BODY}" in main_text

    supp_s2 = [
        p.text
        for p in supplement.paragraphs
        if p.text.startswith("Supplementary Figure S2. Cause-specific rate")
    ]
    assert len(supp_s2) == 1
    assert "moved from the main text" not in supp_s2[0]
    assert "FigureS2_cause_by_approach.png" not in supp_s2[0]

    captions_text = "\n".join(p.text for p in captions.paragraphs)
    for label in (
        "Figure 1.",
        "Figure 2.",
        "Supplementary Figure S1.",
        "Supplementary Figure S2.",
    ):
        assert label in captions_text

    main_md = MAIN_MD.read_text(encoding="utf-8")
    supp_md = SUPP_MD.read_text(encoding="utf-8")
    supp_md_caption = next(
        line
        for line in supp_md.splitlines()
        if line.startswith("**Supplementary Figure S2.** Cause-specific rate")
    )
    assert PHONE in main_md
    assert f"**{S3_LABEL}** {S3_BODY}" in main_md
    assert f"**{S2_LABEL}** {S2_SHORT_BODY}" in main_md
    assert "moved from the main text" not in supp_md_caption
    assert "FigureS2_cause_by_approach.png / .pdf" not in supp_md_caption


def main() -> None:
    update_main_markdown()
    update_main_docx()
    update_supplement_markdown()
    update_supplement_docx()
    create_figure_captions_docx()
    validate_results()
    print("Requested submission edits applied and validated.")


if __name__ == "__main__":
    main()
