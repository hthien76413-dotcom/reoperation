# -*- coding: utf-8 -*-
"""导出与稿件纳入分析完全一致的 450 例队列名单（答辩/审查用）。

字段均为去标识化的研究变量，不含住院号、姓名、入出院日期等可识别信息，
与稿件 Declarations–Data availability 的承诺口径一致（"De-identified
aggregate data..."）——此处精细到患者级，但标识符仅为科研患者编号。

列：科研患者编号 / 性别 / 年龄分层 / 索引手术入路 / 首次手术肠坏死 /
    是否非计划再手术 / 再手术病因 / 十二指肠梗阻机制 / 窗口内竞争事件(确诊死亡)

来源：全部肠旋转不良数据.xlsx（经 _dataprep.py 现算出最终队列）+
      裁定表.xlsx（结局/病因裁定，经 adjudicated()）+
      机制核阅_盲法_答案键.xlsx（十二指肠梗阻的技术缺陷/粘连机制）
"""
import os, sys
import openpyxl
from openpyxl.styles import Font
from _dataprep import load, approach, reop_set, adjudicated, has_necrosis, age_band

D = load()
R = reop_set()
A = adjudicated()
malrot, first = D["malrot"], D["first"]
R = {p for p in R if p in set(malrot)}

MECH_FILE = "机制核阅_盲法_答案键.xlsx"
mech = {}
if os.path.exists(MECH_FILE):
    wk = openpyxl.load_workbook(MECH_FILE, data_only=True)["答案键_评分前勿开"]
    for r in range(2, wk.max_row + 1):
        sid = wk.cell(r, 2).value
        if sid is None:
            continue
        code = str(wk.cell(r, 4).value).strip().upper()
        mech[str(sid).strip()] = {"A": "技术缺陷", "C": "术后粘连"}.get(code, code)

AGE_LABEL = {"Neonate (<28 days)": "新生儿（<28天）",
             "28 days – 1 year": "28天–1岁",
             "≥1 year": "≥1岁"}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "纳入队列(n=450)"

HEADER = ["科研患者编号", "性别", "年龄分层", "索引手术入路", "首次手术肠坏死",
          "是否非计划再手术", "再手术病因", "十二指肠梗阻机制", "窗口内竞争事件(确诊死亡)"]
ws.append(HEADER)
for c in ws[1]:
    c.font = Font(bold=True)

rows = sorted(malrot, key=lambda p: (approach(first[p]), p))
for p in rows:
    op = first[p]
    sex = D["sex"].get(p, "")
    band_en = age_band(D["age_d"].get(p)) if D["age_d"].get(p) is not None else None
    band = AGE_LABEL.get(band_en, "")
    appr = approach(op)
    nec = "是" if has_necrosis(op) else "否"
    is_reop = "是" if p in R else "否"
    cause = A[p]["cause"] if p in R and p in A else ""
    mech_label = mech.get(p, "") if cause == "十二指肠持续梗阻" else ""
    comp = "是" if (p in D["died"] and p not in R) else "否"
    ws.append([p, sex, band, appr, nec, is_reop, cause, mech_label, comp])

widths = [14, 6, 14, 12, 12, 14, 18, 14, 20]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[chr(64 + i)].width = w

ws.freeze_panes = "A2"

OUT = "纳入分析队列名单_450例.xlsx"
wb.save(OUT)
print("saved", OUT, " 共", len(rows), "例")
print("其中非计划再手术", sum(1 for p in rows if p in R), "例")
