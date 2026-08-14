# -*- coding: utf-8 -*-
"""生成【索引手术资格】盲法裁定包：判断每例的索引台次是否为"初次 Ladd 手术"。

盲法设计：只展示索引台次本身，以及该患儿在索引之前的手术记录（判断"既往是否已行
Ladd"所必需）。**索引之后的记录一律不展示**——否则会暴露谁发生了再手术，而候选中
恰有 3 例是本研究的结局病例，知情会污染纳入判断。
病例顺序随机；编号为盲码；答案键单独成表。
"""
import io, sys, random, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import _dataprep as D
# 复用机制核阅那一轮的入路遮蔽函数：资格判定与入路无关，而候选中含唯一的开腹侧
# 十二指肠梗阻病例，若评者看得见入路，纳入/排除的取舍可能差异性地作用于某一臂。
from make_mechanism_blind_review import mask

d = D.load(); ops = d["ops"]; first = d["first"]; coh = set(d["malrot"])
FAR = datetime.datetime(2100, 1, 1)

DIAG = ("活检", "活组织检查", "镜检查", "结肠镜", "胃镜")
THERAP = ("松解", "切除", "吻合", "成形", "修补", "造口", "复位", "探查",
          "引流", "还纳", "减压", "切开", "拉德", "Ladd", "ladd", "根治")
LADDNAME = ("拉德", "Ladd", "ladd")
POSTLADD = ("肠旋转不良术后", "旋转不良(术后)", "旋转不良（术后）", "旋转不良术后")

def diag_only(o):
    sn = o["sname"] or ""
    return any(k in sn for k in DIAG) and not any(k in sn for k in THERAP)

cands = []
for p in sorted(coh):
    o = first[p]; sn = o["sname"] or ""; dx = o["dx"] or ""
    if diag_only(o):
        cat = "①首台仅诊断性"
    elif any(k in dx for k in POSTLADD) or "复发" in dx:
        cat = "②索引为再次/复发 Ladd" if any(k in sn for k in LADDNAME) \
              else "③既往已行 Ladd，本台非 Ladd"
    elif any(k in dx for k in ("术后", "手术的随诊", "手术后随诊")):
        cat = "④诊断栏提及他病术后"
    else:
        continue
    cands.append((p, cat))

print("待裁定 %d 例：" % len(cands))
from collections import Counter
for k, v in sorted(Counter(c for _, c in cands).items()):
    print("   %-24s %d" % (k, v))

rng = random.Random(20260727)
rng.shuffle(cands)

wb = openpyxl.Workbook()
H1 = Font(bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
FILL = PatternFill("solid", fgColor="FFF2CC")

ws = wb.active; ws.title = "0_说明"
for i, line in enumerate([
    "索引手术资格裁定（盲法）",
    "",
    "目的：本研究纳入标准为『接受初次（primary）Ladd 手术的患儿』。下列病例的索引台次",
    "      存疑——或该台仅为诊断性内镜/活检，或诊断栏提示既往已行旋转不良手术。",
    "      请逐例判断该索引台次是否为【初次 Ladd 手术】。",
    "",
    "重要：诊断栏中的『手术的随诊医疗(肠旋转不良术后)』可能是行政随诊编码，不必然代表",
    "      确有前次 Ladd。请以【手术记录本身】为准，不要仅凭诊断栏字符串判定。",
    "",
    "可见信息：索引台次的日期/术名/诊断/【手术记录正文】，以及该患儿在索引之前的手术记录。",
    "          索引之后的记录不予展示（避免暴露结局，候选中含本研究的结局病例）。",
    "          手术记录中的入路措辞（腹腔镜/Troca/气腹等）已遮蔽为 ▉——资格判定与入路无关。",
    "",
    "最有力的证据通常在【手术记录正文】里，例如：",
    "  · 提示既往已行 Ladd：『原切口』『阑尾已切除』『原Ladd's板处粘连』『既往手术瘢痕』",
    "  · 提示本台即为初次 Ladd：首次分离 Ladd's 膜、首次拓宽系膜根、同期附带阑尾切除",
    "",
    "判定选项（填入『判定』列）：",
    "  A 初次 Ladd —— 纳入",
    "  B 非初次·该台仅为诊断性操作（真 Ladd 另有其台或不存在）",
    "  C 非初次·该台是复发/再次 Ladd（既往已行 Ladd）",
    "  D 非初次·既往已行 Ladd，本台并非 Ladd 手术",
    "  E 无法判定 —— 请在备注写明所缺信息",
    "",
    "两位评者独立填写各自表页，完成前请勿互看，也勿打开『答案键』。",
], 1):
    ws.cell(i, 1, line).font = H1 if i == 1 else Font(size=10)
ws.column_dimensions["A"].width = 100

for who in ("评者1", "评者2"):
    s = wb.create_sheet("%s_裁定" % who)
    hdr = ["盲编号", "索引台次日期", "索引台次术名", "索引台次诊断",
           "索引台次·手术记录（术中所见，入路措辞已遮蔽）",
           "索引之前的手术记录", "判定(A–E)", "依据摘句", "备注"]
    for j, h in enumerate(hdr, 1):
        c = s.cell(1, j, h); c.font = H1; c.fill = FILL; c.alignment = WRAP
    for i, (p, cat) in enumerate(cands, 1):
        o = first[p]
        prior = [x for x in ops[p] if x["d"] and o["d"] and x["d"] < o["d"]]
        prior_txt = "\n".join("%s | %s | dx: %s" % (
            x["d"].date(), (x["sname"] or "")[:60], (x["dx"] or "")[:60])
            for x in sorted(prior, key=lambda z: z["d"])) or "（无更早的手术记录）"
        row = [ "IDX%02d" % i, str(o["d"].date()) if o["d"] else "",
                mask(o["sname"] or ""), o["dx"] or "",
                mask(o["narr"] or "") or "（无手术记录正文）",
                prior_txt, "", "", "" ]
        for j, v in enumerate(row, 1):
            c = s.cell(i + 1, j, v); c.alignment = WRAP
        s.row_dimensions[i + 1].height = 150
    for col, w in zip("ABCDEFGHI", (9, 13, 30, 32, 70, 40, 11, 30, 18)):
        s.column_dimensions[col].width = w
    s.freeze_panes = "A2"

k = wb.create_sheet("答案键_评分前勿开")
for j, h in enumerate(["盲编号", "科研患者编号", "算法初判类别", "索引台次日期"], 1):
    k.cell(1, j, h).font = H1
for i, (p, cat) in enumerate(cands, 1):
    k.cell(i + 1, 1, "IDX%02d" % i); k.cell(i + 1, 2, p)
    k.cell(i + 1, 3, cat)
    k.cell(i + 1, 4, str(first[p]["d"].date()) if first[p]["d"] else "")
for col, w in zip("ABCD", (9, 16, 26, 13)):
    k.column_dimensions[col].width = w

OUT = "pending_index_adjudication.xlsx"
wb.save(OUT)
print("\n已生成", OUT, "（%d 例，顺序已随机，索引之后的记录未展示）" % len(cands))
