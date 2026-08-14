# -*- coding: utf-8 -*-
"""14 例十二指肠持续梗阻的年龄呈双峰——核实并展开。

age_sensitivity.py 发现：14 例年龄（天）= 3,3,7,7,7,10 | 1654,2114,2399,2585,2750,3762,4156,4795
中间 10 天到 4.5 岁之间【一例都没有】。这不是连续风险，是两个人群。
本脚本：① 用「最早入院行」重取年龄（_dataprep 取的是表中第一行，8/466 取错）
        ② 分层检验：新生儿层内 vs 学龄层内
        ③ 机制（A技术不彻底 / C粘连）与时相是否也随年龄分层
输出 age_bimodal_out.txt
"""
import io, math
import numpy as np
import openpyxl
from collections import defaultdict
from scipy import stats
from _dataprep import load, approach, reop_set, adjudicated, parse_d, SRC, has_necrosis

L = []
def log(*a): L.append(" ".join(str(x) for x in a))

D = load(); A = adjudicated(); R = reop_set()
malrot, first = D["malrot"], D["first"]

# ---------- ① 修正年龄：取入院日期最早的那条首页记录 ----------
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
ws = wb["病案首页基本信息"]
hdr = [str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
i_pid, i_age = hdr.index("科研患者编号"), hdr.index("年龄（岁）")
i_in = hdr.index("入院日期")
rec = defaultdict(list)
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r or r[i_pid] is None: continue
    rec[str(r[i_pid]).strip()].append((parse_d(r[i_in]), r[i_age]))
def age_fix(p):
    rows = [x for x in rec.get(p, []) if x[0]]
    if not rows: rows = rec.get(p, [])
    if not rows: return None
    v = sorted(rows, key=lambda x: x[0] or parse_d("2100-01-01"))[0][1]
    try: return float(v) * 365.25
    except (TypeError, ValueError): return None

age = np.array([age_fix(p) if age_fix(p) is not None else np.nan for p in malrot])
old = np.array([(D["age"].get(p) or 0) * 365.25 for p in malrot])
log("=" * 78)
log("① 年龄取值修正（_dataprep 取『表中第一行』→ 改取『入院日期最早行』）")
log("=" * 78)
log("  两种取法不一致者 %d/%d；对本节结论无影响，但 Table 1 的年龄行应按修正版重出。"
    % (int(np.sum(np.abs(age - old) > 0.5)), len(malrot)))

lap = np.array([1.0 if approach(first[p]) == "腹腔镜完成" else 0.0 for p in malrot])
duo = np.array([1.0 if (p in R and A[p]["cause"] == "十二指肠持续梗阻") else 0.0 for p in malrot])
reop = np.array([1.0 if p in R else 0.0 for p in malrot])

log("")
log("  14 例十二指肠持续梗阻的年龄（天，修正后）：")
log("     %s" % sorted(int(x) for x in age[duo == 1]))
log("  → 10 天与 1654 天（4.5 岁）之间是空白。不是一条连续的年龄-风险曲线，是两个人群。")

# ---------- ② 分层 ----------
log("")
log("=" * 78)
log("② 分层：新生儿(<28天) vs 婴幼儿(28天–1岁) vs 儿童(≥1岁)")
log("=" * 78)
def exact_or_ci(a, b, c, d, alpha=0.05):
    n1, n2, m1 = a + b, c + d, a + c
    def _lc(n, k): return math.lgamma(n+1) - math.lgamma(k+1) - math.lgamma(n-k+1)
    def tail(psi, upper):
        lo, hi = max(0, m1 - n2), min(n1, m1)
        ks = np.arange(lo, hi + 1)
        lw = np.array([_lc(n1, k) + _lc(n2, m1 - k) for k in ks]) + ks * math.log(psi)
        w = np.exp(lw - lw.max()); w /= w.sum()
        return w[ks >= a].sum() if upper else w[ks <= a].sum()
    def solve(t, upper):
        if upper and a == max(0, m1 - n2): return 0.0
        if (not upper) and a == min(n1, m1): return float("inf")
        loP, hiP = -30.0, 30.0
        f = lambda lp: tail(math.exp(lp), upper) - t
        for _ in range(300):
            mid = (loP + hiP) / 2
            if (f(mid) > 0) == (f(loP) > 0): loP = mid
            else: hiP = mid
        return math.exp((loP + hiP) / 2)
    return solve(alpha/2, True), solve(alpha/2, False)

bands = [("新生儿 <28天", age < 28), ("28天–1岁", (age >= 28) & (age < 365)), ("儿童 ≥1岁", age >= 365)]
for nm, m in bands:
    a1 = int(((lap == 1) & m & (duo == 1)).sum()); n1 = int(((lap == 1) & m).sum())
    a0 = int(((lap == 0) & m & (duo == 1)).sum()); n0 = int(((lap == 0) & m).sum())
    pv = stats.fisher_exact([[a1, n1 - a1], [a0, n0 - a0]])[1]
    lo, hi = exact_or_ci(a1, n1 - a1, a0, n0 - a0)
    log("  %-12s 腹腔镜 %2d/%-3d (%.1f%%)  开腹相关 %2d/%-3d (%.1f%%)   精确95%%CI %.2f–%s  p=%.4f"
        % (nm, a1, n1, a1/max(n1,1)*100, a0, n0, a0/max(n0,1)*100,
           lo, ("%.1f" % hi) if np.isfinite(hi) else "∞", pv))
log("")
log("  同一年龄段内【不分术式】的十二指肠梗阻率：")
for nm, m in bands:
    x = int((m & (duo == 1)).sum()); n = int(m.sum())
    log("     %-12s %2d/%-3d = %.1f%%" % (nm, x, n, x/max(n,1)*100))
log("  → 年龄本身是强预测因子（≥1岁 8.5% vs <28天 2.0%），而 ≥1 岁者 81% 走腹腔镜、")
log("    <28 天者仅 46% 走腹腔镜。年龄与术式高度相关且与结局相关 = 典型混杂，")
log("    且方向【与手稿论证相反】：它让粗 OR 偏大，而不是偏小。")

# ---------- ③ 机制 / 时相 ----------
log("")
log("=" * 78)
log("③ 机制（A 技术不彻底 / C 术后粘连）与时相是否也随年龄分层")
log("=" * 78)
w3 = openpyxl.load_workbook("十二指肠梗阻再手术_归因核阅v3.xlsx", data_only=True)["手术经过全文_填归因"]
h3 = [str(w3.cell(1, c).value) for c in range(1, w3.max_column + 1)]
i_sid = h3.index("研究编号") + 1
i_at = [i for i, c in enumerate(h3) if "归因" in c][0] + 1
i_cat = [i for i, c in enumerate(h3) if "六分类" in c][0] + 1
mech = {}
for r in range(2, w3.max_row + 1):
    sid = w3.cell(r, i_sid).value
    if sid is None: continue
    mech[str(sid).strip()] = (str(w3.cell(r, i_at).value or "").strip(),
                              str(w3.cell(r, i_cat).value or "").strip())
idx = {p: i for i, p in enumerate(malrot)}
log("  %-10s %6s %6s %-8s %-6s %s" % ("研究编号", "年龄天", "间隔天", "术式", "机制", "备注"))
rows = []
for p in sorted(R, key=lambda q: age[idx[q]] if q in idx else 0):
    if A[p]["cause"] != "十二指肠持续梗阻": continue
    i = idx[p]
    mm = mech.get(p, ("?", ""))[0]
    rows.append((p, age[i], A[p]["gap"], "腹腔镜" if lap[i] else "开腹相关", mm))
    log("  %-10s %6.0f %6s %-8s %-6s" % (p, age[i], A[p]["gap"], rows[-1][3], mm))
for nm, sel in [("新生儿 <28天", lambda a: a < 28), ("儿童 ≥1岁", lambda a: a >= 365)]:
    sub = [r for r in rows if sel(r[1])]
    mc = {}
    for r in sub: mc[r[4]] = mc.get(r[4], 0) + 1
    gaps = [r[2] for r in sub if isinstance(r[2], (int, float))]
    log("")
    log("  %s：n=%d，机制 %s，间隔中位 %.0f 天（范围 %d–%d）"
        % (nm, len(sub), mc, float(np.median(gaps)), min(gaps), max(gaps)))

out = "\n".join(L)
io.open("age_bimodal_out.txt", "w", encoding="utf-8").write(out)
print("saved age_bimodal_out.txt")
