# -*- coding: utf-8 -*-
"""按 ICMJE Disclosure Form（短表 docx 版）为每位作者各生成一份填好的表。

模板 `ICMJE_Disclosure_Form_template.docx` 是 ICMJE 通用短表（当初从 Surgical
Endoscopy 投稿系统下载，正文不含刊名，任何期刊通用）。Springer 的 Ped Surg Int
只硬性要求正文里的 Declarations 段（Funding / Competing interests），**未见明文
要求上传 ICMJE 表**；备一份是为投稿系统临时索要，不是必交件。

填写口径与稿件 Declarations 完全一致：无任何资助、无任何利益冲突，故 13 项
一律填 None。**这只是把作者已在稿件中作出的声明誊到表上，不是替作者判断**——
第 10 项（学会/委员会任职，有偿无偿均需申报）之类容易漏，每位作者务必自查后
再签名。

产出：ICMJE_<姓名拼音>.docx ×5（已 gitignore，可复现）

用法：  python3 make_icmje_forms.py
"""
import copy
import io
import os
import re

from docx import Document

TEMPLATE = "ICMJE_Disclosure_Form_template.docx"
SRC = "PedSurgInt_manuscript_v1.md"

# 签署日期。投稿日若不是当天，作者自行改；留空亦可。
SIGN_DATE = "11 September 2026"
MS_NUMBER = ""          # 尚未投出，无稿号；投稿系统若已给号，填在这里重跑

# 与稿件扉页署名顺序一致
AUTHORS = ["Jun Shu", "Kai Zheng", "Hongqiang Bian", "Jun Yang", "Xin Wang"]

# 13 项披露一律为无。措辞用 "None"，与表头 "or indicate none" 对齐。
NONE = "None"

# 表 1 中「Name all entities…」列的列下标，以及 13 个条目所在的行下标
ENTITY_COL = 3
ITEM_ROWS = [2] + list(range(4, 16))     # 第 1 项在 r2，第 2–13 项在 r4–r15
CERT_ROW = 18                            # "I certify that…" 所在行


def read_title(path=SRC):
    """从主稿扉页取题名，避免表里的题名与稿件不一致。"""
    text = io.open(path, encoding="utf-8").read()
    m = re.search(r"^\*\*Title:\*\*\s*(.+)$", text, flags=re.M)
    if not m:
        raise SystemExit("在 %s 里找不到 **Title:** 行" % path)
    return m.group(1).strip()


def set_cell(cell, text, bold=None):
    """把单元格内容换成 text，沿用原有字体（保留第一个 run 的格式）。"""
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        run = p.runs[0]
    else:
        run = p.add_run(text)
    if bold is not None:
        run.bold = bold


def append_to_label(cell, value):
    """在 "Date:"、"Your Name:" 这类标签后面接上内容，标签本身不动。"""
    p = cell.paragraphs[0]
    run = copy.deepcopy(p.runs[-1]._element) if p.runs else None
    if run is None:
        p.add_run(" " + value)
        return
    p._p.append(run)
    p.runs[-1].text = " " + value
    p.runs[-1].bold = True


def fill(author, title):
    doc = Document(TEMPLATE)
    head, body = doc.tables[0], doc.tables[1]

    append_to_label(head.rows[1].cells[0], SIGN_DATE)
    append_to_label(head.rows[2].cells[0], author)
    append_to_label(head.rows[3].cells[0], title)
    if MS_NUMBER:
        append_to_label(head.rows[4].cells[0], MS_NUMBER)

    for ri in ITEM_ROWS:
        set_cell(body.rows[ri].cells[ENTITY_COL], NONE)

    # 「place an "X" next to the following statement」——表里没有独立复选框，
    # 整行是一个合并单元格，故把 X 加在句子前面，问题措辞一字未动。
    cert = body.rows[CERT_ROW].cells[0].paragraphs[0]
    lead = copy.deepcopy(cert.runs[0]._element)
    cert._p.insert(list(cert._p).index(cert.runs[0]._element), lead)
    cert.runs[0].text = "X   "
    cert.runs[0].bold = True

    out = "ICMJE_%s.docx" % author.replace(" ", "_")
    doc.save(out)
    return out


def main():
    if not os.path.exists(TEMPLATE):
        raise SystemExit("缺少模板 %s" % TEMPLATE)
    title = read_title()
    print("题名（取自 %s 扉页）：\n  %s\n" % (SRC, title))
    for a in AUTHORS:
        print("  已生成", fill(a, title))
    print("\n13 项披露均填 None，与稿件 Declarations 一致。"
          "\n每位作者签名前请自查第 10 项（学会/委员会任职，有偿无偿均须申报）"
          "与第 2 项（任何来源的科研经费）。")


if __name__ == "__main__":
    main()
