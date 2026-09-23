# -*- coding: utf-8 -*-
"""核心结论二（再手术病因构成因术式而异）三线表 -> docx。

数字均取自 原始数据分析代码/结论二_再手术病因构成因术式而异/step1-4 的实测输出，
与稿件 Table 2 完全一致，未手写。
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROWS = [
    ("1. 十二指肠持续梗阻", "12/255", "0/195", "4.7 vs. 0.0", "NE（≥2.19）", "0.0016", "0.010"),
    ("2. 坏死/穿孔/吻合口并发症", "2/255", "6/195", "0.8 vs. 3.1", "0.25（0.02–1.42）", "0.0821", "0.411"),
    ("3. 粘连性肠梗阻（非十二指肠）", "4/255", "1/195", "1.6 vs. 0.5", "3.09（0.30–153.03）", "0.3944", "1.000"),
    ("4. 合并畸形漏诊", "3/255", "1/195", "1.2 vs. 0.5", "2.31（0.18–121.86）", "0.6367", "1.000"),
    ("5. 肠扭转复发", "2/255", "0/195", "0.8 vs. 0.0", "NE（≥0.14）", "0.5078", "1.000"),
    ("6. 其他", "1/255", "4/195", "0.4 vs. 2.1", "0.19（0.00–1.93）", "0.1713", "0.685"),
]
HEADER = ["病因（单选）", "腹腔镜 n/N", "开腹相关 n/N", "发生率(%) 腹腔镜 vs. 开腹相关",
          "OR（精确 95% CI）", "Fisher p", "Holm 校正 p"]
TITLE = "表 2　再手术病因构成的组间比较（核心结论二）"
NOTE = ("注：以全队列为分母（腹腔镜完成 255 例、开腹相关 195 例）；OR 为条件精确"
        "（Cornfield）95% 置信区间；Holm 逐步法对六个病因族内校正。持续性十二指肠梗阻"
        "在开腹相关组无一例发生，为零单元，OR 不可估（NE），仅单侧下界有意义，"
        "风险差为可解释的效应量：+4.7 个百分点（Newcombe 95% CI 1.9 至 8.0）；"
        "该项是六个病因中唯一经 Holm 校正后仍 <0.05 者。数据来源：本院肠旋转不良"
        "患儿数据库 450 例。")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, spec in kwargs.items():
        tag = "w:%s" % edge
        el = tcBorders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tcBorders.append(el)
        el.set(qn("w:sz"), str(spec.get("sz", 4)))
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:color"), spec.get("color", "000000"))


def no_border(cell):
    set_cell_border(cell, top={"sz": 0, "val": "nil"}, bottom={"sz": 0, "val": "nil"},
                     left={"sz": 0, "val": "nil"}, right={"sz": 0, "val": "nil"})


doc = Document()
for st in doc.styles:
    if st.name == "Normal":
        st.font.name = "Times New Roman"
        st.font.size = Pt(10.5)
        rPr = st.element.get_or_add_rPr()
        ea = rPr.find(qn("w:rFonts"))
        if ea is None:
            ea = OxmlElement("w:rFonts")
            rPr.append(ea)
        ea.set(qn("w:eastAsia"), "宋体")

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE)
r.bold = True
r.font.size = Pt(11)
r.font.name = "Times New Roman"
rPr = r._element.get_or_add_rPr()
ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "黑体"); rPr.append(ea)

ncols = len(HEADER)
table = doc.add_table(rows=1 + len(ROWS), cols=ncols)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True

widths = [Cm(4.6), Cm(2.4), Cm(2.6), Cm(3.6), Cm(3.6), Cm(2.0), Cm(2.2)]
for i, w in enumerate(widths):
    for cell in table.columns[i].cells:
        cell.width = w

for j, text in enumerate(HEADER):
    cell = table.cell(0, j)
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "黑体"); rPr.append(ea)
    no_border(cell)

for i, row in enumerate(ROWS, start=1):
    bold_row = row[-1] != "1.000" and float(row[-1]) < 0.05
    for j, text in enumerate(row):
        cell = table.cell(i, j)
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.bold = bold_row
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        rPr = run._element.get_or_add_rPr()
        ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "宋体"); rPr.append(ea)
        no_border(cell)

# 三线表边框：表格顶线、表头下线、表格底线，其余一律无线
for j in range(ncols):
    set_cell_border(table.cell(0, j), top={"sz": 12, "val": "single"})
    set_cell_border(table.cell(0, j), bottom={"sz": 8, "val": "single"})
    set_cell_border(table.cell(len(ROWS), j), bottom={"sz": 12, "val": "single"})

note_p = doc.add_paragraph()
note_run = note_p.add_run(NOTE)
note_run.font.size = Pt(9)
note_run.italic = True
note_run.font.name = "Times New Roman"
rPr = note_run._element.get_or_add_rPr()
ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "宋体"); rPr.append(ea)

doc.save("核心结论二_三线表.docx")
print("saved 核心结论二_三线表.docx")
