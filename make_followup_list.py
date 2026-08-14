# -*- coding: utf-8 -*-
"""导出 42 天窗口内"自动出院/放弃治疗"患儿的随访工作表。

用途：这些患儿在竞争风险分析中被计为竞争事件，其合法性取决于"离院后确已死亡、
不可能在外院再次手术"。本表供电话随访逐例确认，并留出填写栏。
"""
import io, sys, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import _dataprep as D

d = D.load(); coh = set(d["malrot"]); first = d["first"]
adj = D.adjudicated(); ad = d["age_d"]
AB, DE = d["abandon"], d["death"]

# 原始表里取住院号与出院记录原文
wb0 = openpyxl.load_workbook(D.SRC, read_only=True, data_only=True)
ws_head = wb0["病案首页基本信息"]
hdr = [str(c) for c in next(ws_head.iter_rows(min_row=1, max_row=1, values_only=True))]
i_pid, i_zyh = hdr.index("科研患者编号"), hdr.index("住院号")
i_in, i_out = hdr.index("入院日期"), hdr.index("出院日期")
zyh = {}
for r in ws_head.iter_rows(min_row=2, values_only=True):
    if not r or r[i_pid] is None: continue
    p = str(r[i_pid]).strip()
    zyh.setdefault(p, []).append((D.parse_d(r[i_in]), D.parse_d(r[i_out]),
                                  str(r[i_zyh]).strip() if r[i_zyh] else ""))

ws_dis = wb0["住院病历出院记录"]
dh = [str(c) for c in next(ws_dis.iter_rows(min_row=1, max_row=1, values_only=True))]
j_pid, j_out, j_cond = dh.index("科研患者编号"), dh.index("出院诊断"), dh.index("出院情况")
disrec = {}
for r in ws_dis.iter_rows(min_row=2, values_only=True):
    if not r or r[j_pid] is None: continue
    p = str(r[j_pid]).strip()
    disrec.setdefault(p, []).append((str(r[j_out] or ""), str(r[j_cond] or "")))

KW = ["放弃治疗", "自动出院", "自行出院"]
def flag_sentence(p):
    """摘出触发'自动出院/放弃'判定的原文片段。"""
    for dx, cond in disrec.get(p, []):
        blob = dx + " " + cond
        for k in KW:
            i = blob.find(k)
            if i >= 0:
                return blob[max(0, i-40):i+40].replace("\n", " ")
    return "（出院记录中未直接检出关键词，请核对病案）"

def gap(p):
    dis, idx = d["last_dis"].get(p), first[p]["d"]
    return (dis - idx).days if (dis and idx) else None

rows = []
for p in sorted(coh & AB):
    g = gap(p)
    if g is None or not (0 <= g <= 42): continue
    o = first[p]
    zs = sorted(zyh.get(p, []), key=lambda x: (x[0] or datetime.datetime(2100,1,1)))
    rows.append(dict(
        pid=p,
        zyh_all="；".join(sorted({z[2] for z in zs if z[2]})),
        sex=d["sex"].get(p, ""),
        age_d=round(ad[p], 1) if ad.get(p) is not None else "",
        idx_date=o["d"].date() if o["d"] else "",
        approach=D.approach(o),
        sname=(o["sname"] or "")[:70],
        dx=(o["dx"] or "")[:70],
        necrosis="是" if D.has_necrosis(o) else "否",
        last_dis=d["last_dis"][p].date() if d["last_dis"].get(p) else "",
        gap=g,
        also_death="是" if p in DE else "否",
        reop="是" if p in adj else "否",
        quote=flag_sentence(p)))

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "随访工作表"
H1 = Font(bold=True, size=10); WRAP = Alignment(wrap_text=True, vertical="top")
FILL = PatternFill("solid", fgColor="DDEBF7"); FILL2 = PatternFill("solid", fgColor="FFF2CC")

cols = [("序号",6),("科研患者编号",14),("住院号（全部住院）",22),("性别",6),
        ("索引手术日龄(天)",12),("索引手术日期",13),("入路",11),
        ("索引手术名称",34),("术中/首台诊断",34),("索引时肠坏死",10),
        ("末次出院日期",13),("距索引手术(天)",11),("院内已记录死亡",11),
        ("42天内本院再手术",13),("触发'自动出院/放弃'的原文片段",42),
        ("——以下由随访填写——",20),
        ("随访日期",12),("是否接通",10),
        ("结局（已死亡/存活/失访）",16),("死亡日期",12),
        ("离院后是否曾在其他医院手术",18),("何院/何时/何术式",26),
        ("随访者",10),("备注",26)]
for j,(h,w) in enumerate(cols,1):
    c=ws.cell(1,j,h); c.font=H1; c.alignment=WRAP
    c.fill = FILL2 if j>=16 else FILL
    ws.column_dimensions[get_column_letter(j)].width=w

for i,r in enumerate(rows,1):
    vals=[i,r["pid"],r["zyh_all"],r["sex"],r["age_d"],str(r["idx_date"]),r["approach"],
          r["sname"],r["dx"],r["necrosis"],str(r["last_dis"]),r["gap"],
          r["also_death"],r["reop"],r["quote"],"","","","","","","","",""]
    for j,v in enumerate(vals,1):
        c=ws.cell(i+1,j,v); c.alignment=WRAP
    ws.row_dimensions[i+1].height=54
ws.freeze_panes="C2"

n = wb.create_sheet("说明")
for i,l in enumerate([
    "自动出院／放弃治疗患儿电话随访工作表",
    "",
    "背景：本研究把 42 天窗口内的死亡作为竞争事件。竞争风险分析成立的前提是",
    "      『该事件确实终止了再次手术的可能』。院内死亡满足这一前提；",
    "      『自动出院／放弃治疗』则不然——患儿可能在其他医院接受了再次手术，",
    "      而本院登记无法捕获。故须逐例电话确认。",
    "",
    "本表共 %d 例，即队列（450 例）中 42 天窗口内出院方式含" % len(rows),
    "『放弃治疗／自动出院／自行出院』者。其中已在院内记录为死亡者见『院内已记录死亡』列。",
    "",
    "填写要点：",
    "  · 『结局』只填三选一：已死亡 / 存活 / 失访。",
    "  · 若『存活』，务必追问『离院后是否曾在其他医院手术』——这是本次随访的核心问题，",
    "    一旦为『是』，该例就不能再作为竞争事件，需要改按结局或删失处理。",
    "  · 『失访』也要如实填写，不要留空；失访例数需要写进论文的局限。",
    "  · 原始资料无姓名与联系电话，请凭『住院号』在院内系统调取。",
    "",
    "回填后把本文件发回，即可据以更新分析与正文。",
], 1):
    n.cell(i,1,l).font = Font(bold=(i==1), size=10)
n.column_dimensions["A"].width=92

OUT="自动出院随访工作表.xlsx"
wb.save(OUT)
print("已导出 %s：%d 例" % (OUT, len(rows)))
from collections import Counter
print("入路分布:", dict(Counter(r["approach"] for r in rows)))
print("院内已记录死亡:", sum(1 for r in rows if r["also_death"]=="是"), "例")
print("42 天内本院再手术:", sum(1 for r in rows if r["reop"]=="是"), "例")
