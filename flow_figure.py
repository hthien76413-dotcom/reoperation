# -*- coding: utf-8 -*-
"""Figure 1：参与者纳入 / 分组流程图（STROBE participant flow）。

2026-07-26 修订：原图残留 v1（22 例名单）时代的内容——总数硬编码为 4.5%、
结局写成 "for obstruction"、脚注还在讲 "22 reoperations total"。现全部按裁定表
重算，并按审稿意见补上【结局裁定的排除账目及其在两组间的分布】。
"""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import openpyxl
from _dataprep import load, approach, reop_set, is_malrot, ADJ as ADJ_PATH

D = load()
ops, first, malrot = D["ops"], D["first"], D["malrot"]
n_surg = len(ops)                         # 有手术记录的数据库病例
n_cohort = len(malrot)                    # 450

# 2026-07-28 发现：Excluded 框写 n=49，但逐项列出的六类只加总到 33，缺口 16（实为18，
# 因原 33 本身也算错，见下）——真正原因是 49 例排除分两种性质，此前混在一个框里且
# 只写了其中一种：(a) 索引台次根本不是恶转不良适应证手术 31 例；(b) 索引台次文本
# 命中恶转不良关键词、但被第四遍资格裁定判为非首台 Ladd 18 例（见 NONPRIMARY_REASON）。
# 31 例的六类逐例核对如下（2026-07-28 手工阅片分类，凭 dx+sname 文本），
# 断言其覆盖 is_malrot()==False 的全集且不重不漏，防止队列再变动时又悄悄脱节。
EXCL_CAT = {
    "congenital duodenal obstruction":  {"10535230", "197410", "35806888", "9081394", "9980960"},
    "diaphragmatic hernia / eventration": {"1260608", "4590144", "4748594", "535755", "8048425"},
    "other intestinal atresia":         {"10189692", "13606882", "2563281"},
    "appendectomy":                     {"18597467", "54940", "4219878"},
    "adhesiolysis / volvulus, not coded": {"35785789", "5887078", "8977706"},
    "other conditions":                 {"1953361", "20399813", "2603818", "3564649", "3590321",
                                          "36167351", "3788090", "4158180", "4240068", "4321080",
                                          "4408103", "5304860"},
}
_not_malrot = {p for p in ops if not is_malrot(first[p])}
_covered = set().union(*EXCL_CAT.values())
assert _covered == _not_malrot, (
    "★ EXCL_CAT 与当前 is_malrot()==False 的病例集不一致（缺 %s，多 %s）——"
    "队列筛选逻辑或数据已变动，请重新阅片分类，不要想当然改数字。"
    % (_not_malrot - _covered, _covered - _not_malrot))
grp = {"lap": [], "conv": [], "open": []}
for p in malrot:
    a = approach(first[p])
    grp["lap" if a == "腹腔镜完成" else "conv" if a == "中转开腹" else "open"].append(p)
REOP = reop_set()
def reop(ps): return sum(1 for p in ps if p in REOP)
n_lap, n_conv, n_open = len(grp["lap"]), len(grp["conv"]), len(grp["open"])
r_lap, r_conv, r_open = reop(grp["lap"]), reop(grp["conv"]), reop(grp["open"])
n_excl = n_surg - n_cohort

# 结局裁定账目：53 例候选 → 39 纳入；排除项按首台入路分组
ws = openpyxl.load_workbook(ADJ_PATH, data_only=True)["A_结局裁定"]
COH = set(malrot)
adj = {}
for r in range(2, ws.max_row + 1):
    pid = ws.cell(r, 2).value
    if pid is None: continue
    adj.setdefault(str(ws.cell(r, 12).value or "").strip(), []).append(str(pid).strip())
n_cand = sum(len(v) for v in adj.values())
def split(status):
    ps = adj.get(status, [])
    lap = sum(1 for p in ps if p in COH and approach(first[p]) == "腹腔镜完成")
    opr = sum(1 for p in ps if p in COH and approach(first[p]) != "腹腔镜完成")
    return len(ps), lap, opr, len(ps) - lap - opr
n_plan, plan_lap, plan_opr, plan_out = split("排除-计划性")
n_late = split("排除-晚期")[0]
n_miss = split("排除-资料缺失")[0]
# 裁定表标"纳入"的是 39 例，其中 3 例的患儿在第四遍资格裁定中被判非首台 Ladd、
# 已不在 450 例队列内，故不能计为队列结局。必须按队列收缩，否则右框写 39
# 而三个入路框合计 36，图内自相矛盾（2026-07-28 发现）。
_inc_all = split("纳入")[0]
n_inc = sum(1 for p in adj.get("纳入", []) if p in COH)
n_nonprim = _inc_all - n_inc

fig, ax = plt.subplots(figsize=(9.6, 7.4))
ax.set_xlim(0, 12); ax.set_ylim(0, 10); ax.axis("off")
CX = 4.0          # 主流程中轴（左移，给右侧两个排除框让位）

def box(x, y, w, h, text, fc="#F4F6F8", ec="#33475B", fs=9.3, ha="center"):
    ax.add_patch(FancyBboxPatch((x - w/2, y - h/2), w, h,
        boxstyle="round,pad=0.02,rounding_size=0.08",
        linewidth=1.3, edgecolor=ec, facecolor=fc, zorder=2))
    tx = x if ha == "center" else x - w/2 + 0.12
    ax.text(tx, y, text, ha=ha, va="center", fontsize=fs, zorder=3, linespacing=1.4)

def connector(x1, y1, x2, y2):
    """Plain connector segment used for clean orthogonal branches."""
    ax.plot([x1, x2], [y1, y2], color="#33475B", lw=1.35,
            solid_capstyle="round", zorder=1)

def arrow(x1, y1, x2, y2):
    """Single, high-contrast arrowhead at the destination only."""
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
        mutation_scale=17, lw=1.35, color="#33475B",
        shrinkA=0, shrinkB=0, zorder=1))

# A 顶部
box(CX, 9.2, 7.0, 1.0,
    f"Children undergoing surgery in the single-center\nintestinal-malrotation database (Dec 2012 – Jun 2026)\n(n = {n_surg})")
# 排除（右侧，左对齐列表）——两种不同性质的排除理由分两段列，理由见上方注释
NR = D["nonprimary_reason"]
excl_lines = "\n".join(f"   • {k}  {len(v)}" for k, v in EXCL_CAT.items())
box(9.75, 7.30, 4.1, 4.60,
    f"Excluded (n = {n_excl})\n"
    f"Index operation not a malrotation-\nindication procedure ({len(_not_malrot)}):\n"
    f"{excl_lines}\n"
    f"Index operation matched malrotation-\nrelated terms but was not a primary\n"
    f"Ladd procedure, per blinded\nadjudication ({len(D['nonprimary'])}, §2.4):\n"
    f"   • diagnostic endoscopy or biopsy only  {len(NR['B_诊断性内镜或活检'])}\n"
    f"   • redo Ladd for recurrent malrotation  {len(NR['C_复发或再次Ladd'])}\n"
    f"   • previous Ladd already performed  {len(NR['D_既往已行Ladd'])}",
    fc="#FBEEEE", ec="#B0413E", fs=7.6, ha="left")
arrow(CX, 8.68, CX, 7.52)
arrow(CX, 7.95, 7.66, 7.95)
# B 队列
box(CX, 7.0, 6.0, 0.95,
    f"Primary malrotation-indication (Ladd) cohort\n(n = {n_cohort})", fc="#EAF1F8", ec="#2C6FBB")
connector(CX, 6.50, CX, 5.98)
connector(1.3, 5.98, 6.7, 5.98)
ax.text(CX, 6.24, "Classified by index surgical approach",
        ha="center", fontsize=8.6, style="italic", color="#555",
        bbox=dict(facecolor="white", edgecolor="none", pad=1.5), zorder=2)

# C 三组
xs = [1.3, 4.0, 6.7]
labels = [f"Laparoscopic\ncompletion\n(n = {n_lap})",
          f"Conversion\nto open\n(n = {n_conv})",
          f"Open\n(n = {n_open})"]
for x, lb in zip(xs, labels):
    box(x, 5.15, 2.5, 1.15, lb, fc="#F4F6F8")
    arrow(x, 5.98, x, 5.75)
    arrow(x, 4.55, x, 3.75)
# D 事件
events = [f"Unplanned reoperation\nwithin 42 days\n{r_lap} / {n_lap} ({r_lap/n_lap*100:.1f}%)",
          f"Unplanned reoperation\nwithin 42 days\n{r_conv} / {n_conv} ({r_conv/n_conv*100:.1f}%)",
          f"Unplanned reoperation\nwithin 42 days\n{r_open} / {n_open} ({r_open/n_open*100:.1f}%)"]
for x, ev in zip(xs, events):
    box(x, 3.15, 2.5, 1.15, ev, fc="#FDF6E9", ec="#C08A00")

# 结局裁定账目（右侧，与事件行齐平）
box(10.10, 3.35, 3.6, 2.3,
    f"Candidate reoperations adjudicated (n = {n_cand})\n"
    f"Excluded by blinded adjudication:\n"
    f"   • planned staged procedure  {n_plan}\n"
    f"        ({plan_lap} laparoscopic, {plan_opr} open-related,\n"
    f"         {plan_out} outside the cohort)\n"
    f"   • beyond the 42-day window  {n_late}\n"
    f"   • records unavailable  {n_miss}\n"
    f"   • index operation not a primary Ladd  {n_nonprim}\n"
    f"→ Unplanned reoperations included: {n_inc}",
    fc="#FBEEEE", ec="#B0413E", fs=8.0, ha="left")
arrow(7.97, 3.35, 8.27, 3.35)

# 总事件 + 脚注
n_tot = r_lap + r_conv + r_open
ax.text(CX, 1.75, f"Total unplanned reoperations: {n_tot} / {n_cohort} ({n_tot/n_cohort*100:.1f}%)",
        ha="center", fontsize=9.6, fontweight="bold")
ax.text(CX, 1.05,
        "Planned staged procedures were excluded irrespective of timing. Counting them as outcomes,\n"
        "or reverting the fourth-pass eligibility exclusions above, did not change the\n"
        "comparison between approaches (Table 4).",
        ha="center", fontsize=8.0, style="italic", color="#555", linespacing=1.3)

fig.tight_layout()
fig.savefig("Figure1_flow_EN.png", dpi=300)
fig.savefig("Figure1_flow_EN.pdf")
# ps.fonttype=42 让 TrueType 字体嵌入 EPS，满足官方「Vector graphics containing
# fonts must have the fonts embedded」。必须放在 PDF 保存【之后】：本机 matplotlib
# 在该参数生效时写 PDF 会抛 ValueError: bytes must be in range(0, 256)
# （字体子集化的共用代码路径所致），只影响 PDF，不影响 EPS 本身。
matplotlib.rcParams["ps.fonttype"] = 42
fig.savefig("Figure1_flow_EN.eps")   # 矢量投稿版
print(f"n_surg={n_surg}  cohort={n_cohort}  lap={n_lap}/{r_lap}  conv={n_conv}/{r_conv}  open={n_open}/{r_open}")
print("saved Figure1_flow_EN.png / .pdf")
