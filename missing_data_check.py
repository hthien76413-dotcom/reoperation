# -*- coding: utf-8 -*-
"""逐变量核查 466 例队列中每个入稿变量的缺失/不可判定情况，
给 Methods 2.5 / Limitations 的缺失数据声明提供实测依据（而非估计）。"""
import io
import _dataprep
from _dataprep import load, approach, reop_set
from datetime import datetime

D = load()
malrot, first, ops = D["malrot"], D["first"], D["ops"]

L = []
def log(*a): L.append(" ".join(str(x) for x in a))

n = len(malrot)
log(f"队列 n={n}")
log("")

# 1) 性别 / 年龄（病案首页，Table 1 用）
miss_sex = sum(1 for p in malrot if not D["sex"].get(p))
miss_age = sum(1 for p in malrot if D["age"].get(p) is None)
log(f"[Table 1] 性别缺失: {miss_sex}/{n}")
log(f"[Table 1] 年龄缺失: {miss_age}/{n}")

# 2) 首次手术分类所需文本（sname+narr 全空 => 只能靠默认规则归类为"开腹"，存在误判风险）
empty_op_text = sum(1 for p in malrot if not first[p]["sname"].strip() and not first[p]["narr"].strip())
log(f"[Methods 2.3 索引术式分类] 首次手术记录 手术名称+手术经过 均为空文本: {empty_op_text}/{n}  "
    f"(若为空则默认归类为'开腹'，存在误判风险)")

# 3) 首次手术日期缺失（CIF/间隔计算依赖 first[p]['d']）
miss_date = sum(1 for p in malrot if first[p]["d"] is None)
log(f"[竞争风险主分析] 首次手术日期缺失: {miss_date}/{n}  (缺失则无法计入CIF时间轴)")

# 4) 新生儿判定（Apgar代理）——本身是二元代理定义，不存在"缺失"，只有"无法判定为新生儿"
neo_n = sum(1 for p in malrot if p in D["neo"])
log(f"[新生儿代理] 有Apgar记录: {neo_n}/{n}；无Apgar记录（按非新生儿处理）: {n-neo_n}/{n}  "
    f"(设计上是代理定义，非缺失，但已在Limitations说明其局限)")

# 5) 出院记录用于死亡/放弃判定与竞争事件时间（last_dis）
no_dis_record = sum(1 for p in malrot if D["last_dis"].get(p) is None)
log(f"[竞争事件时间] 无任何可用出院日期记录: {no_dis_record}/{n}")

# 死亡/放弃者中，出院日期缺失导致按90天行政删失处理（见 strengthen.py 逻辑）的例数
death_abandon = (D["death"] | D["abandon"]) & set(malrot)
da_no_date = sum(1 for p in death_abandon if D["last_dis"].get(p) is None)
log(f"[竞争风险] 死亡/放弃者中缺出院日期（因而按90天行政删失处理而非计为竞争事件）: "
    f"{da_no_date}/{len(death_abandon)}")

# 6) 坏死/畸形关键词抽取依赖的诊断文本（出院记录+病理诊断）是否完全为空
import openpyxl
from collections import defaultdict
wb = openpyxl.load_workbook(_dataprep.SRC, read_only=True, data_only=True)
dx = defaultdict(str)
ws = wb["住院病历出院记录"]
for r in ws.iter_rows(min_row=2, values_only=True):
    if r and r[0] is not None:
        dx[str(r[0]).strip()] += " " + " ".join(str(r[i]) for i in (3, 5) if len(r) > i and r[i])
ws = wb["病案病理诊断"]
for r in ws.iter_rows(min_row=2, values_only=True):
    if r and r[0] is not None and len(r) > 2 and r[2]:
        dx[str(r[0]).strip()] += " " + str(r[2])
no_dx_text = sum(1 for p in malrot if not dx.get(p, "").strip())
log(f"[Table 1 畸形/坏死关键词抽取] 出院诊断+病理诊断文本完全缺失: {no_dx_text}/{n}")

out = "\n".join(L)
io.open("missing_data_check_out.txt", "w", encoding="utf-8").write(out)
try: print(out)
except UnicodeEncodeError: print(out.encode("gbk", "replace").decode("gbk"))
