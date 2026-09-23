# -*- coding: utf-8 -*-
"""核心结论一（总体再手术率两种术式无差异）三线表 -> docx。

数字均取自 原始数据分析代码/结论一_总体再手术率两种术式无差异/step1-4 的实测输出，
与稿件 Table 3 上半部分（总体率一行 + 敏感性/校正行）一致，未手写。
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROWS = [
    ("腹腔镜完成", "24/255", "9.4", "6.4–13.6", "—", "—"),
    ("中转开腹", "6/53", "11.3", "5.3–22.6", "—", "—"),
    ("开腹", "6/142", "4.2", "2.0–8.9", "—", "—"),
    ("全队列", "36/450", "8.0", "5.8–10.9", "—", "—"),
    ("粗比值比（腹腔镜 vs. 开腹相关）", "—", "—", "—", "OR 1.58（0.74–3.6）", "0.224"),
    ("风险差（腹腔镜 − 开腹相关）", "—", "—", "—", "+3.3 个百分点（−2.0 至 +8.2）", "—"),
    ("Firth，未校正", "—", "—", "—", "OR 1.55（0.76–3.16）", "0.215"),
    ("Firth ＋新生儿 ＋首次手术坏死", "—", "—", "—", "OR 1.64（0.75–3.60）", "0.209"),
    ("Firth ＋年龄段 ＋首次手术坏死", "—", "—", "—", "OR 1.59（0.72–3.49）", "0.246"),
]
HEADER = ["分组 / 分析", "再手术 n/N", "发生率 (%)", "95% CI (%)", "效应量 (95% CI)", "P 值"]
TITLE = "表 1　总体非计划再手术率的组间比较（核心结论一）"
NOTE = ("注：腹腔镜完成 vs. 开腹相关（中转开腹 + 开腹）为预设主要比较，Fisher 精确检验，"
        "OR 为条件精确（Cornfield）区间，风险差为 Newcombe 区间；下方三行为 Firth 惩罚似然"
        "逻辑回归的敏感性/校正分析。数据来源：本院肠旋转不良患儿数据库 450 例。")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, spec in kwargs.items():
        tag = "w:%s" % edge
        el = tcBorders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tcBorders.append(el)
        el.set(qn("w:sz"), str(spec.get("sz", 4)))
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:color"), spec.get("color", "000000"))


def no_border(cell):
    set_cell_border(cell, top={"sz": 0, "val": "nil"}, bottom={"sz": 0, "val": "nil"},
                     left={"sz": 0, "val": "nil"}, right={"sz": 0, "val": "nil"})


doc = Document()
sec = doc.sections[0]
for st in doc.styles:
    if st.name == "Normal":
        st.font.name = "Times New Roman"
        st.font.size = Pt(10.5)
        rPr = st.element.get_or_add_rPr()
        ea = rPr.find(qn("w:rFonts"))
        if ea is None:
            ea = OxmlElement("w:rFonts")
            rPr.append(ea)
        ea.set(qn("w:eastAsia"), "宋体")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE)
r.bold = True
r.font.size = Pt(11)
r.font.name = "Times New Roman"
rPr = r._element.get_or_add_rPr()
ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "黑体"); rPr.append(ea)

ncols = len(HEADER)
table = doc.add_table(rows=1 + len(ROWS), cols=ncols)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True

widths = [Cm(5.2), Cm(2.2), Cm(2.2), Cm(2.4), Cm(4.6), Cm(2.0)]
for i, w in enumerate(widths):
    for cell in table.columns[i].cells:
        cell.width = w

for j, text in enumerate(HEADER):
    cell = table.cell(0, j)
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "黑体"); rPr.append(ea)
    no_border(cell)

for i, row in enumerate(ROWS, start=1):
    for j, text in enumerate(row):
        cell = table.cell(i, j)
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.font.size = Pt(10.5)
        run.font.name = "Times New Roman"
        rPr = run._element.get_or_add_rPr()
        ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "宋体"); rPr.append(ea)
        no_border(cell)

# 三线表边框：表格顶线、表头下线、表格底线，其余一律无线
for j in range(ncols):
    set_cell_border(table.cell(0, j), top={"sz": 12, "val": "single"})
    set_cell_border(table.cell(0, j), bottom={"sz": 8, "val": "single"})
    set_cell_border(table.cell(len(ROWS), j), bottom={"sz": 12, "val": "single"})

note_p = doc.add_paragraph()
note_run = note_p.add_run(NOTE)
note_run.font.size = Pt(9)
note_run.italic = True
note_run.font.name = "Times New Roman"
rPr = note_run._element.get_or_add_rPr()
ea = OxmlElement("w:rFonts"); ea.set(qn("w:eastAsia"), "宋体"); rPr.append(ea)

doc.save("核心结论一_三线表.docx")
print("saved 核心结论一_三线表.docx")
