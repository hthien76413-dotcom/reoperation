# -*- coding: utf-8 -*-
"""生成《39例再手术患儿详细临床资料.xlsx》

数据源：
  1) 全部肠旋转不良数据.xlsx  —— 主库（去标识化），覆盖全部39例
  2) 25例再次手术（多次住院）.xlsx —— 补充 姓名 / 手术医生 / 确切二开原因（9例可匹配）
  3) 17例丢失的再次手术病例原始临床资料.xlsx —— 补录主库缺失的再手术记录（pid 4042304）
  4) 裁定表.xlsx —— 病例名单、再手术日期、间隔、共识原因（唯一权威）
"""
import openpyxl, io, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict
from datetime import datetime
from _dataprep import load, approach, reop_set, adjudicated, has_necrosis, has_resection

DIR = r"D:\全部肠旋转不良\③肠旋转不良术后再手术"
SRC = r"D:\全部肠旋转不良\全部肠旋转不良数据.xlsx"
NEW25 = DIR + r"\25例再次手术（多次住院）.xlsx"
LOST17 = DIR + r"\17例丢失的再次手术病例原始临床资料.xlsx"
OUT = DIR + r"\39例再手术患儿详细临床资料.xlsx"

def s(v):
    return "" if v is None else str(v).strip()

def dstr(d):
    if isinstance(d, datetime): return d.strftime("%Y-%m-%d")
    return s(d)[:10]

D = load(); ADJ = adjudicated(); REOP = sorted(reop_set())
first, ops = D["first"], D["ops"]

# ---------------- 主库读取 ----------------
wb0 = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def sheet_rows(name):
    ws = wb0[name]
    it = ws.iter_rows(values_only=True)
    hdr = [s(c) for c in next(it)]
    return hdr, it

# 病案首页：pid -> [每次住院]
hdr, it = sheet_rows("病案首页基本信息")
ix = {k: hdr.index(k) for k in hdr}
adm = defaultdict(list)
for r in it:
    if not r or r[0] is None: continue
    adm[s(r[0])].append({h: r[i] for h, i in ix.items()})

# 手术记录：pid -> [手术]
hdr, it = sheet_rows("住院病历手术记录")
surg = defaultdict(list)
for r in it:
    if not r or r[0] is None: continue
    pid, vid, sdate, dx, sname, anes, narr = (list(r) + [None]*7)[:7]
    d = sdate if isinstance(sdate, datetime) else None
    if d is None and sdate:
        try: d = datetime.strptime(s(sdate)[:10], "%Y-%m-%d")
        except: d = None
    surg[s(pid)].append(dict(vid=s(vid), d=d, dx=s(dx), sname=s(sname),
                             anes=s(anes), narr=s(narr)))

# 出院记录：pid -> [记录]
hdr, it = sheet_rows("住院病历出院记录")
disc = defaultdict(list)
for r in it:
    if not r or r[0] is None: continue
    disc[s(r[0])].append(dict(vid=s(r[1]), ru=s(r[2]), rudx=s(r[3]),
                              jing=s(r[4]), chudx=s(r[5]), chu=s(r[6])))

# 入院记录（新生儿科 / 儿科）
neo_rec = {}
hdr, it = sheet_rows("新生儿科入院记录")
ixn = {k: hdr.index(k) for k in hdr}
for r in it:
    if not r or r[0] is None: continue
    p = s(r[0])
    if p in neo_rec: continue
    neo_rec[p] = {h: s(r[i]) for h, i in ixn.items()}
ped_rec = {}
hdr, it = sheet_rows("儿科入院记录")
ixp = {k: hdr.index(k) for k in hdr}
for r in it:
    if not r or r[0] is None: continue
    p = s(r[0])
    if p in ped_rec: continue
    ped_rec[p] = {h: s(r[i]) for h, i in ixp.items()}

# 影像 / 病理
def load_reports(sheet, cols):
    hdr, it = sheet_rows(sheet)
    ii = {k: hdr.index(k) for k in cols if k in hdr}
    out = defaultdict(list)
    for r in it:
        if not r or r[0] is None: continue
        out[s(r[0])].append({k: s(r[i]) for k, i in ii.items()})
    return out
us  = load_reports("超声报告", ["超声报告名称","超声检查时间","超声检查所见","超声检查结论"])
xr  = load_reports("X线报告", ["报告名称","检查时间","检查所见","检查结论"])
ct  = load_reports("CT报告", ["报告名称","检查时间","检查所见","检查结论"])
pa  = load_reports("病理报告", ["病理报告名称","病理检查时间","肉眼所见","病理诊断"])

# ---------------- 可选数据源：姓名 / 手术医生 / 确切二开原因 ----------------
# 《25例再次手术（多次住院）.xlsx》为可选补充源；该表是"多次住院"切片（含大量晚期再手术），
# 与本39例仅少数可匹配。文件缺失时跳过，相关三列留空，不影响其余内容。
import os
zy_info = {}      # 住院号 -> dict(name, surgeon, reason)
if os.path.exists(NEW25):
    wbn = openpyxl.load_workbook(NEW25, data_only=True)
    ws = wbn["二次手术患儿统计"]
    cur_name = ""
    for r in range(2, ws.max_row + 1):
        nm = s(ws.cell(r, 2).value)
        if nm and not nm.startswith("（"): cur_name = nm
        zy = s(ws.cell(r, 4).value)
        if zy:
            zy_info.setdefault(zy, {})["name"] = cur_name
    ws = wbn["非新生儿"]
    cur_name = ""
    for r in range(2, ws.max_row + 1):
        nm = s(ws.cell(r, 3).value)
        if nm: cur_name = nm
        zy = s(ws.cell(r, 5).value)
        if zy:
            d = zy_info.setdefault(zy, {})
            d.setdefault("name", cur_name)
            if s(ws.cell(r, 13).value): d["surgeon"] = s(ws.cell(r, 13).value)
            if s(ws.cell(r, 14).value): d["reason"] = s(ws.cell(r, 14).value)
    print("[信息] 已载入可选源 25例再次手术（多次住院）.xlsx")
else:
    print("[提示] 未找到《25例再次手术（多次住院）.xlsx》，姓名/手术医生/确切二开原因三列将留空")

# pid -> 该患儿全部住院号
pid_zys = {p: [s(a.get("住院号")) for a in adm.get(p, []) if s(a.get("住院号"))] for p in REOP}

def from_new(p, key):
    for zy in pid_zys.get(p, []):
        v = zy_info.get(zy, {}).get(key)
        if v: return v
    return ""

# ---------------- 补录：主库缺失的手术记录（统一走 _supplements） ----------------
from _supplements import load_supplements
SUPP = load_supplements()
SUPP_REOP = {p: dict(d=s["d"], sname=s["sname"], dx=s["dx"], anes=s["anes"],
                     narr=s["narr"], note="【补录】" + s["source"])
             for p, s in SUPP.items() if s["kind"] == "reop"}
SUPP_INDEX = {p: dict(d=s["d"], sname=s["sname"], dx=s["dx"], anes=s["anes"],
                      narr=s["narr"], note="【补录】" + s["source"])
              for p, s in SUPP.items() if s["kind"] == "index"}

# ---------------- 组装每例 ----------------
def pick_ops(p):
    vs = sorted(surg.get(p, []), key=lambda x: (x["d"] or datetime(2100, 1, 1)))
    if p in SUPP_INDEX:
        vs = [v for v in vs if v["d"] != SUPP_INDEX[p]["d"]]
        vs = sorted(vs + [SUPP_INDEX[p]], key=lambda x: (x["d"] or datetime(2100, 1, 1)))
    idx = vs[0] if vs else None
    dr = ADJ[p]["d_reop"]
    re_ = None
    if dr:
        cand = [v for v in vs[1:] if v["d"]]
        if cand:
            re_ = min(cand, key=lambda v: abs((v["d"] - dr).days))
            if abs((re_["d"] - dr).days) > 7: re_ = None
    if re_ is None and p in SUPP_REOP:
        re_ = SUPP_REOP[p]
    return idx, re_

def in_window(tstr, lo, hi):
    if not tstr: return False
    m = re.search(r"(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})", tstr)
    if not m: return False
    try: t = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except: return False
    return lo <= t <= hi

cases = []
for p in REOP:
    a = ADJ[p]
    idx, re_ = pick_ops(p)
    a0 = adm.get(p, [{}])[0]
    rec = neo_rec.get(p) or ped_rec.get(p) or {}
    cases.append(dict(
        pid=p, zy=s(a["zyh"]), name=from_new(p, "name"),
        surgeon=from_new(p, "surgeon"), reason_new=from_new(p, "reason"),
        sex=s(a0.get("性别")), age=a0.get("年龄（岁）"),
        bw=s(a0.get("新生儿入院体重（g）")),
        neo="是" if p in D["neo"] else "否",
        apgar1=rec.get("Apgar评分1分钟", ""), apgar5=rec.get("Apgar评分5分钟", ""),
        chief=rec.get("主诉", ""), hpi=rec.get("现病史", ""), initdx=rec.get("初步诊断", ""),
        grp=approach(first[p]),
        nec="是" if has_necrosis(first[p]) else "否",
        res="是" if has_resection(first[p]) else "否",
        d_index=dstr(a["d_index"]), d_reop=dstr(a["d_reop"]), gap=a["gap"], cause=a["cause"],
        idx_op=idx or {}, re_op=re_ or {},
        disc=disc.get(p, []),
        us=us.get(p, []), xr=xr.get(p, []), ct=ct.get(p, []), pa=pa.get(p, []),
    ))

# ---------------- 写 Excel ----------------
HF = PatternFill("solid", fgColor="1F3864"); HFONT = Font(bold=True, color="FFFFFF", size=10)
F_LAP = PatternFill("solid", fgColor="FDEBD0"); F_OPN = PatternFill("solid", fgColor="E8F0FE")
thin = Side(style="thin", color="BBBBBB"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = openpyxl.Workbook()

def mkheader(ws, cols, widths, row=1):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row, j, c); cell.fill = HF; cell.font = HFONT; cell.border = BORD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row + 1, 1)

# ===== Sheet1 总表 =====
ws1 = wb.active; ws1.title = "总表"
cols = ["序号","研究编号","住院号(首台)","姓名","性别","年龄(岁)","新生儿","出生体重(g)",
        "首台日期","首台入路","首台手术名称","首台术中诊断","首台坏死","首台肠切除",
        "再手术日期","间隔(天)","再手术原因(裁定)","再手术手术名称","再手术术中诊断",
        "手术医生","确切二开原因(源:25例表)"]
w = [5,11,12,10,6,8,7,10,11,10,30,26,7,8,11,7,18,32,26,16,40]
mkheader(ws1, cols, w)
for i, c in enumerate(cases, 2):
    vals = [i-1, c["pid"], c["zy"], c["name"], c["sex"], c["age"], c["neo"], c["bw"],
            c["d_index"], c["grp"], c["idx_op"].get("sname",""), c["idx_op"].get("dx",""),
            c["nec"], c["res"], c["d_reop"], c["gap"], c["cause"],
            c["re_op"].get("sname",""), c["re_op"].get("dx",""), c["surgeon"], c["reason_new"]]
    fill = F_LAP if c["grp"] == "腹腔镜完成" else F_OPN
    for j, v in enumerate(vals, 1):
        cell = ws1.cell(i, j, v); cell.border = BORD; cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=(j in (11,12,18,19,21)))
    ws1.row_dimensions[i].height = 42

# ===== Sheet2/3 手术记录全文 =====
for title, key in [("首台手术记录", "idx_op"), ("再手术记录", "re_op")]:
    ws = wb.create_sheet(title)
    mkheader(ws, ["序号","研究编号","住院号","姓名","日期","手术名称","术中诊断","麻醉","手术经过（原文）","备注"],
             [5,11,12,10,11,32,28,10,120,26])
    for i, c in enumerate(cases, 2):
        o = c[key]
        vals = [i-1, c["pid"], c["zy"], c["name"], dstr(o.get("d")), o.get("sname",""),
                o.get("dx",""), o.get("anes",""), o.get("narr",""), o.get("note","")]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(i, j, v); cell.border = BORD
            cell.alignment = Alignment(vertical="top", wrap_text=(j in (6,7,9,10)))
        ws.row_dimensions[i].height = 120

# ===== Sheet4 入院记录 =====
ws4 = wb.create_sheet("入院记录")
mkheader(ws4, ["序号","研究编号","住院号","姓名","Apgar1min","Apgar5min","主诉","现病史","初步诊断"],
         [5,11,12,10,10,10,30,90,40])
for i, c in enumerate(cases, 2):
    vals = [i-1, c["pid"], c["zy"], c["name"], c["apgar1"], c["apgar5"], c["chief"], c["hpi"], c["initdx"]]
    for j, v in enumerate(vals, 1):
        cell = ws4.cell(i, j, v); cell.border = BORD
        cell.alignment = Alignment(vertical="top", wrap_text=(j >= 7))
    ws4.row_dimensions[i].height = 100

# ===== Sheet5 出院记录（每次住院一行） =====
ws5 = wb.create_sheet("出院记录")
mkheader(ws5, ["序号","研究编号","住院号","姓名","入院情况","入院诊断","诊疗经过","出院诊断","出院情况"],
         [5,11,12,10,40,30,80,40,40])
rr = 2
for i, c in enumerate(cases, 1):
    for d in c["disc"]:
        vals = [i, c["pid"], c["zy"], c["name"], d["ru"], d["rudx"], d["jing"], d["chudx"], d["chu"]]
        for j, v in enumerate(vals, 1):
            cell = ws5.cell(rr, j, v); cell.border = BORD
            cell.alignment = Alignment(vertical="top", wrap_text=(j >= 5))
        ws5.row_dimensions[rr].height = 90
        rr += 1

# ===== Sheet6 围手术期影像/病理（首台前7天 ~ 再手术后14天） =====
ws6 = wb.create_sheet("影像与病理")
mkheader(ws6, ["序号","研究编号","住院号","姓名","类别","检查名称","检查时间","所见","结论/诊断"],
         [5,11,12,10,8,24,14,80,50])
rr = 2
for i, c in enumerate(cases, 1):
    lo = ADJ[c["pid"]]["d_index"]; hi = ADJ[c["pid"]]["d_reop"]
    if not (lo and hi): continue
    lo = datetime(lo.year, lo.month, lo.day); hi = datetime(hi.year, hi.month, hi.day)
    lo = lo.replace(day=lo.day) ;
    from datetime import timedelta
    lo2, hi2 = lo - timedelta(days=7), hi + timedelta(days=14)
    for tag, lst, kn, kt, kf, kc in [
        ("超声", c["us"], "超声报告名称","超声检查时间","超声检查所见","超声检查结论"),
        ("X线", c["xr"], "报告名称","检查时间","检查所见","检查结论"),
        ("CT",  c["ct"], "报告名称","检查时间","检查所见","检查结论"),
        ("病理", c["pa"], "病理报告名称","病理检查时间","肉眼所见","病理诊断")]:
        for rec in lst:
            if not in_window(rec.get(kt,""), lo2, hi2): continue
            vals = [i, c["pid"], c["zy"], c["name"], tag, rec.get(kn,""), rec.get(kt,"")[:16],
                    rec.get(kf,""), rec.get(kc,"")]
            for j, v in enumerate(vals, 1):
                cell = ws6.cell(rr, j, v); cell.border = BORD
                cell.alignment = Alignment(vertical="top", wrap_text=(j >= 8))
            ws6.row_dimensions[rr].height = 70
            rr += 1

# ===== Sheet7 说明 =====
ws7 = wb.create_sheet("数据说明")
notes = [
    ("《39例再手术患儿详细临床资料》", True),
    ("", False),
    ("【病例范围】", True),
    ("39例 = 裁定表.xlsx Sheet『A_结局裁定』中『是否纳入』=纳入 的全部病例。", False),
    ("结局定义：所有【计划外】再手术，发生于首台术后≤42天内或与首台同属一次住院。", False),
    ("已排除：计划性分期手术11例、晚期2例、资料缺失1例（共53例候选）。", False),
    ("", False),
    ("【再手术原因分类（6类，2026-07-26定稿）】", True),
    ("1 十二指肠持续梗阻（不含内在畸形） / 2 肠坏死·穿孔·吻合口并发症 / 3 粘连性肠梗阻 /", False),
    ("4 合并畸形漏诊 / 5 肠扭转复发 / 6 其他", False),
    ("第4类『合并畸形漏诊』= 再手术时才发现或证实的合并消化道畸形（十二指肠膜·蹼·隔膜、环状胰腺、肠闭锁等）；", False),
    ("  首台已知的畸形不计入本类。凡再手术发现内在十二指肠畸形者一律归第4类，不再计入第1类。", False),
    ("沿革：原『造口相关』并入『其他』；『复发肠扭转/redo-Ladd』更名『肠扭转复发』。", False),
    ("分类经两位评者盲法独立裁定后取共识；评者原始副本保留评分当时的类别名称，不予回填。", False),
    ("", False),
    ("【数据来源】", True),
    ("1) 全部肠旋转不良数据.xlsx —— 主库（去标识化），覆盖全部39例的手术/出入院/影像/病理记录。", False),
    ("2) 25例再次手术（多次住院）.xlsx —— 【可选源】补充『姓名』『手术医生』『确切二开原因』；", False),
    ("   该表为『多次住院』切片（含大量晚期再手术），与本39例仅少数可匹配；若该文件不存在则三列全部留空。", False),
    ("3) 17例丢失的再次手术病例原始临床资料.xlsx —— 补录主库缺失的再手术记录（研究编号4042304，", False),
    ("   住院号1041483，第二次住院记录未导入主库；该行『备注』列已标注来源）。", False),
    ("4) 裁定表.xlsx —— 病例名单、首台/再手术日期、间隔天数、共识原因（唯一权威来源）。", False),
    ("", False),
    ("【字段说明】", True),
    ("首台入路：腹腔镜完成 / 中转开腹 / 开腹，来自手术记录文本判定并经盲法裁定。", False),
    ("新生儿：以入院记录含Apgar评分为代理（非严格生后日龄）。", False),
    ("首台坏死 / 首台肠切除：自首台手术记录文本提取（已排除『未见坏死』等否定表述）。", False),
    ("影像与病理：仅列出首台术前7天至再手术后14天窗口内的检查。", False),
    ("", False),
    ("【颜色】总表中 橙色=首台腹腔镜完成，蓝色=首台开腹相关（中转+开腹）。", False),
]
for i, (t, bold) in enumerate(notes, 1):
    cell = ws7.cell(i, 1, t)
    cell.font = Font(bold=bold, size=12 if i == 1 else 10)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
ws7.column_dimensions["A"].width = 120

wb.save(OUT)
print("saved", OUT)
print("总表行数 =", len(cases))
print("有姓名 =", sum(1 for c in cases if c["name"]),
      "| 有手术医生 =", sum(1 for c in cases if c["surgeon"]),
      "| 有确切二开原因 =", sum(1 for c in cases if c["reason_new"]))
print("再手术记录缺失 =", sum(1 for c in cases if not c["re_op"].get("sname")))
